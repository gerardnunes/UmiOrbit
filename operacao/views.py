import io

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db.models import Count, Q, Prefetch
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views import View
from django.views.generic import (
    CreateView,
    DetailView,
    ListView,
    TemplateView,
    UpdateView,
)

from .forms import (
    AtualizacaoForm,
    ChecklistForm,
    ChecklistImportForm,
    DemandaForm,
    FiltroDemandaForm,
    MissaoForm,
    PerfilForm,
    UsuarioForm,
)
from .models import (
    STATUS_ABERTOS,
    Atualizacao,
    Checklist,
    ChecklistImportacao,
    ChecklistItem,
    ChecklistItemEvento,
    ChecklistItemFoto,
    Demanda,
    Missao,
    Perfil,
    Prioridade,
    Sistema,
    Status,
    StatusMissao,
    Tipo,
)
from .permissions import TrabalhadorRequiredMixin
from .services import ler_excel, aplicar_import

User = get_user_model()


# =========================================================
# CONSTANTES
# =========================================================
GRAVES = (Prioridade.CRITICA, Prioridade.ALTA)

ESTADOS = {
    "NAO_APTO": ("Não apto para operar", "Há pendências que impedem a operação."),
    "ATENCAO": ("Apto com restrições", "Nada impede a saída, mas há pendências graves."),
    "APTO": ("Apto para operar", "Nenhuma pendência grave em aberto."),
}

# Mapeamento: nome do sistema da DEMANDA -> local no barco (JS)
MAPA_SISTEMA_LOCAL = {
    "Motor": "motor", "Propulsão": "helice", "Bateria": "bateria",
    "Baterias": "bateria", "Starlink": "starlink", "Comunicação": "antena",
    "Casco": "casco", "Estrutura": "casco", "Mastro": "mastro",
    "Leme": "leme", "Cabine": "cabine", "Inversor": "inversor",
}
LOCAL_FALLBACK = "cabine"


# =========================================================
# HELPERS
# =========================================================
def _nivel(bloqueantes, graves):
    return "NAO_APTO" if bloqueantes else "ATENCAO" if graves else "APTO"


def _registrar(demanda, autor, texto="", anterior="", novo=""):
    return Atualizacao.objects.create(
        demanda=demanda, autor=autor, texto=texto,
        status_anterior=anterior, status_novo=novo,
    )


def _evento(item, tipo, autor=None, texto=""):
    return ChecklistItemEvento.objects.create(
        item=item, tipo=tipo, autor=autor, texto=texto,
    )


# =========================================================
# DASHBOARD
# =========================================================
class DashboardView(LoginRequiredMixin, TemplateView):
    template_name = "dashboard.html"

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)

        # ---------- Demandas ----------
        ativas = Demanda.objects.exclude(status=Status.CANCELADA)
        abertas = ativas.filter(status__in=STATUS_ABERTOS)

        total = ativas.count()
        concluidas = ativas.filter(status=Status.CONCLUIDA).count()
        bloqueantes = abertas.filter(bloqueia_operacao=True).count()
        graves = abertas.filter(prioridade__in=GRAVES).count()
        nivel = _nivel(bloqueantes, graves)

        # ---------- Sistemas / embarcações ----------
        sistemas = (
            Sistema.objects
            .annotate(
                total=Count("demandas", filter=~Q(demandas__status=Status.CANCELADA)),
                feitas=Count("demandas", filter=Q(demandas__status=Status.CONCLUIDA)),
                abertas=Count("demandas", filter=Q(demandas__status__in=STATUS_ABERTOS)),
                bloqueantes=Count("demandas", filter=Q(
                    demandas__status__in=STATUS_ABERTOS,
                    demandas__bloqueia_operacao=True,
                )),
                graves=Count("demandas", filter=Q(
                    demandas__status__in=STATUS_ABERTOS,
                    demandas__prioridade__in=GRAVES,
                )),
            )
        )
        for s in sistemas:
            s.pct = s.feitas * 100 // s.total if s.total else 100
            s.nivel = _nivel(s.bloqueantes, s.graves)
            s.nivel_label = ESTADOS[s.nivel][0]

        # ---------- Missões ativas ----------
        missoes_ativas = (
            Missao.objects
            .filter(status__in=[
                StatusMissao.PREPARACAO,
                StatusMissao.EM_OPERACAO,
                StatusMissao.PAUSADA,
            ])
            .select_related("responsavel")
            .prefetch_related("checklists__itens")
            .order_by("-criado_em")[:6]
        )
        for m in missoes_ativas:
            cl = list(m.checklists.all())
            if cl:
                m.total_itens = sum(c.itens.count() for c in cl)
                m.itens_feitos = sum(
                    c.itens.filter(concluido=True).count() for c in cl
                )
                m.progresso = (
                    m.itens_feitos * 100 // m.total_itens
                    if m.total_itens else 0
                )
            else:
                m.total_itens = 0
                m.itens_feitos = 0
                m.progresso = 0

        # ---------- Checklists ativos ----------
        checklists_ativos = (
            Checklist.objects
            .exclude(status=Checklist.Status.ARQUIVADO)
            .select_related("missao", "criado_por")
            .prefetch_related("itens")
            .order_by("-atualizado_em")[:6]
        )
        for c in checklists_ativos:
            c.total_itens = c.itens.count()
            c.itens_feitos = c.itens.filter(concluido=True).count()
            c.pct_calc = (
                c.itens_feitos * 100 // c.total_itens if c.total_itens else 0
            )

        total_checklists = Checklist.objects.count()
        checklists_concluidos = Checklist.objects.filter(
            status=Checklist.Status.CONCLUIDO
        ).count()

        # ---------- Auditoria recente ----------
        fotos_recentes = (
            ChecklistItemFoto.objects
            .select_related("item", "item__checklist", "enviado_por")
            .order_by("-enviado_em")[:8]
        )

        # ---------- Diário de bordo ----------
        diario = []
        for a in Atualizacao.objects.select_related("demanda", "autor")[:6]:
            diario.append({
                "tipo": "demanda",
                "quando": a.criado_em,
                "titulo": a.demanda.titulo,
                "texto": a.get_status_novo_display()
                         or (a.texto[:60] if a.texto else "—"),
                "autor": a.autor.get_full_name() if a.autor else "",
                "icone": "🔧",
                "cor": "demanda",
            })
        for ev in ChecklistItemEvento.objects.select_related(
            "item", "item__checklist", "autor"
        )[:6]:
            diario.append({
                "tipo": "checklist",
                "quando": ev.criado_em,
                "titulo": f"{ev.item.checklist.titulo} → {ev.item.titulo}",
                "texto": ev.get_tipo_display(),
                "autor": ev.autor.get_full_name() if ev.autor else "",
                "icone": "☑",
                "cor": "checklist",
            })
        diario.sort(key=lambda x: x["quando"], reverse=True)
        diario = diario[:10]

        # ---------- Problemas no barco (demandas reais) ----------
        problemas_barco = []
        usados = set()
        alertas = (
            abertas
            .filter(Q(bloqueia_operacao=True) | Q(prioridade__in=GRAVES))
            .select_related("sistema")
            .order_by("-bloqueia_operacao", "prioridade", "prazo")
        )
        for d in alertas:
            nome = d.sistema.nome if d.sistema else ""
            local = MAPA_SISTEMA_LOCAL.get(nome, LOCAL_FALLBACK)
            if local in usados:
                continue
            usados.add(local)
            problemas_barco.append({
                "local": local,
                "titulo": d.titulo,
                "codigo": d.codigo,
                "prioridade": d.get_prioridade_display(),
                "bloqueante": bool(d.bloqueia_operacao),
                "sistema": nome,
            })

        # ---------- Contexto ----------
        ctx.update(
            # demanda
            estado=nivel,
            estado_titulo=ESTADOS[nivel][0],
            estado_texto=ESTADOS[nivel][1],
            prontidao=concluidas * 100 // total if total else 100,
            total=total,
            concluidas=concluidas,
            n_abertas=abertas.count(),
            n_andamento=abertas.filter(status=Status.ANDAMENTO).count(),
            n_pendentes=abertas.filter(status=Status.PENDENTE).count(),
            n_bloqueadas=abertas.filter(status=Status.BLOQUEADA).count(),
            n_bloqueantes=bloqueantes,
            n_atrasadas=abertas.filter(prazo__lt=timezone.localdate()).count(),
            n_criticas=abertas.filter(prioridade=Prioridade.CRITICA).count(),

            # sistemas
            sistemas=sistemas,

            # missões
            missoes_ativas=missoes_ativas,
            total_missoes=Missao.objects.count(),
            missoes_em_operacao=Missao.objects.filter(
                status=StatusMissao.EM_OPERACAO
            ).count(),

            # checklists
            checklists_ativos=checklists_ativos,
            total_checklists=total_checklists,
            checklists_concluidos=checklists_concluidos,

            # auditoria
            fotos_recentes=fotos_recentes,
            total_fotos=ChecklistItemFoto.objects.count(),

            # diário
            diario=diario,

            # barco
            problemas_barco=problemas_barco,

            # hora
            agora=timezone.now(),
        )
        return ctx


# =========================================================
# DEMANDAS
# =========================================================
class DemandaListView(LoginRequiredMixin, ListView):
    template_name = "demanda_list.html"
    context_object_name = "demandas"
    paginate_by = 20

    def get_filtro(self):
        dados = self.request.GET.copy()
        if "status" not in dados:
            dados["status"] = "ABERTAS"
        return FiltroDemandaForm(dados)

    def get_queryset(self):
        qs = Demanda.objects.select_related("sistema", "responsavel")
        f = self.get_filtro()
        if f.is_valid():
            d = f.cleaned_data
            if d["q"]:
                termo = d["q"].strip()
                cond = (
                    Q(titulo__icontains=termo)
                    | Q(item_faltante__icontains=termo)
                    | Q(descricao__icontains=termo)
                )
                if termo.upper().startswith("L42-") and termo[4:].isdigit():
                    cond |= Q(pk=int(termo[4:]))
                qs = qs.filter(cond)
            if d["status"] == "ABERTAS":
                qs = qs.filter(status__in=STATUS_ABERTOS)
            elif d["status"]:
                qs = qs.filter(status=d["status"])
            for campo in ("prioridade", "tipo", "sistema", "responsavel"):
                if d[campo]:
                    qs = qs.filter(**{campo: d[campo]})
            if d["atrasadas"]:
                qs = qs.filter(
                    status__in=STATUS_ABERTOS,
                    prazo__lt=timezone.localdate(),
                )
        return qs.ordenadas()

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["filtro"] = self.get_filtro()
        return ctx


class DemandaDetailView(LoginRequiredMixin, DetailView):
    model = Demanda
    template_name = "demanda_detail.html"
    context_object_name = "demanda"

    def get_queryset(self):
        return Demanda.objects.select_related(
            "sistema", "responsavel", "criado_por"
        )

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["form"] = AtualizacaoForm()
        ctx["atualizacoes"] = self.object.atualizacoes.select_related("autor")
        return ctx


class DemandaCreateView(TrabalhadorRequiredMixin, CreateView):
    model = Demanda
    form_class = DemandaForm
    template_name = "demanda_form.html"

    def get_initial(self):
        initial = super().get_initial()
        if self.request.GET.get("tipo") in Tipo.values:
            initial["tipo"] = self.request.GET["tipo"]
        return initial

    def form_valid(self, form):
        form.instance.criado_por = self.request.user
        resp = super().form_valid(form)
        _registrar(self.object, self.request.user, "Demanda criada.")
        messages.success(self.request, f"{self.object.codigo} registrada.")
        return resp


class DemandaUpdateView(TrabalhadorRequiredMixin, UpdateView):
    model = Demanda
    form_class = DemandaForm
    template_name = "demanda_form.html"

    def form_valid(self, form):
        anterior = form.initial.get("status", "")
        mudou_status = "status" in form.changed_data
        resp = super().form_valid(form)
        if mudou_status:
            _registrar(
                self.object, self.request.user,
                anterior=anterior, novo=self.object.status,
            )
        elif form.changed_data:
            _registrar(self.object, self.request.user, "Demanda editada.")
        messages.success(self.request, f"{self.object.codigo} atualizada.")
        return resp


class AssumirDemandaView(TrabalhadorRequiredMixin, View):
    def post(self, request, pk):
        demanda = get_object_or_404(Demanda, pk=pk)
        anterior = demanda.status
        demanda.responsavel = request.user
        if demanda.status == Status.PENDENTE:
            demanda.status = Status.ANDAMENTO
        demanda.save()
        novo = demanda.status if demanda.status != anterior else ""
        _registrar(
            demanda, request.user, "Assumiu a demanda.",
            anterior if novo else "", novo,
        )
        messages.success(request, f"Você assumiu {demanda.codigo}.")
        return redirect(demanda)


class AtualizarDemandaView(TrabalhadorRequiredMixin, View):
    def post(self, request, pk):
        demanda = get_object_or_404(Demanda, pk=pk)
        form = AtualizacaoForm(request.POST)
        if not form.is_valid():
            messages.error(request, "Escreva um comentário ou escolha um novo status.")
            return redirect(demanda)
        novo = form.cleaned_data["status"]
        texto = form.cleaned_data["texto"].strip()
        anterior = demanda.status
        if novo and novo != anterior:
            demanda.status = novo
            if novo == Status.ANDAMENTO and not demanda.responsavel:
                demanda.responsavel = request.user
            demanda.save()
        else:
            novo = ""
        _registrar(demanda, request.user, texto, anterior if novo else "", novo)
        messages.success(request, "Atualização registrada.")
        return redirect(demanda)


# =========================================================
# EQUIPE / PERFIL
# =========================================================
class EquipeView(LoginRequiredMixin, TemplateView):
    template_name = "equipe.html"

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        usuarios = (
            User.objects
            .filter(is_active=True)
            .select_related("perfil")
            .annotate(
                n_abertas=Count(
                    "demandas",
                    filter=Q(demandas__status__in=STATUS_ABERTOS),
                ),
                n_feitas=Count(
                    "demandas",
                    filter=Q(demandas__status=Status.CONCLUIDA),
                ),
            )
            .order_by("first_name", "username")
        )
        ctx["diretores"] = [
            u for u in usuarios
            if getattr(u, "perfil", None)
            and u.perfil.papel == Perfil.Papel.DIRETOR
        ]
        ctx["trabalhadores"] = [
            u for u in usuarios
            if not getattr(u, "perfil", None)
            or u.perfil.papel != Perfil.Papel.DIRETOR
        ]
        return ctx


class MeuPerfilView(LoginRequiredMixin, View):
    template_name = "perfil.html"

    def _render(self, request, uform, pform):
        minhas = (
            Demanda.objects
            .filter(responsavel=request.user, status__in=STATUS_ABERTOS)
            .select_related("sistema")
            .ordenadas()
        )
        return render(request, self.template_name, {
            "uform": uform, "pform": pform, "minhas": minhas,
        })

    def get(self, request):
        return self._render(
            request,
            UsuarioForm(instance=request.user),
            PerfilForm(instance=request.user.perfil),
        )

    def post(self, request):
        uform = UsuarioForm(request.POST, instance=request.user)
        pform = PerfilForm(request.POST, instance=request.user.perfil)
        if uform.is_valid() and pform.is_valid():
            uform.save()
            pform.save()
            messages.success(request, "Perfil atualizado.")
            return redirect("meu_perfil")
        return self._render(request, uform, pform)


# =========================================================
# MANUTENÇÃO
# =========================================================
class ManutencaoView(LoginRequiredMixin, ListView):
    model = Demanda
    template_name = "manutencao.html"
    context_object_name = "manutencoes"

    def get_queryset(self):
        return (
            Demanda.objects
            .filter(tipo=Tipo.MANUTENCAO)
            .select_related("sistema", "responsavel")
            .ordenadas()
        )

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        qs = self.get_queryset()
        ctx["total"] = qs.count()
        ctx["abertas"] = qs.filter(status__in=STATUS_ABERTOS).count()
        ctx["bloqueantes"] = qs.filter(
            status__in=STATUS_ABERTOS, bloqueia_operacao=True
        ).count()
        ctx["concluidas"] = qs.filter(status=Status.CONCLUIDA).count()
        return ctx


# =========================================================
# MISSÕES
# =========================================================
class MissaoListView(LoginRequiredMixin, ListView):
    model = Missao
    template_name = "missao_list.html"
    context_object_name = "missoes"

    def get_queryset(self):
        return (
            Missao.objects
            .select_related("responsavel")
            .prefetch_related("checklist")
            .order_by("-criado_em")
        )

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["missao_atual"] = (
            Missao.objects
            .filter(status__in=[
                StatusMissao.PREPARACAO,
                StatusMissao.EM_OPERACAO,
                StatusMissao.PAUSADA,
            ])
            .select_related("responsavel")
            .prefetch_related("checklist")
            .order_by("-criado_em")
            .first()
        )
        ctx["status_missao"] = StatusMissao
        return ctx


class MissaoCreateView(LoginRequiredMixin, CreateView):
    model = Missao
    form_class = MissaoForm
    template_name = "missao_form.html"

    def form_valid(self, form):
        form.instance.criado_por = self.request.user
        messages.success(self.request, "Missão criada.")
        return super().form_valid(form)

    def get_success_url(self):
        return reverse("missoes")


# =========================================================
# CHECKLISTS
# =========================================================
class ChecklistListView(LoginRequiredMixin, ListView):
    model = Checklist
    template_name = "checklist_list.html"
    context_object_name = "checklists"

    def get_queryset(self):
        qs = (
            Checklist.objects
            .select_related("missao", "criado_por")
            .prefetch_related("itens")
            .order_by("-criado_em")
        )
        status = self.request.GET.get("status")
        if status:
            qs = qs.filter(status=status)
        return qs

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        checklists = list(ctx["checklists"])

        grupos = {}
        for c in checklists:
            key = c.missao_id
            if key not in grupos:
                grupos[key] = {"missao": c.missao, "checklists": []}
            grupos[key]["checklists"].append(c)

        com_missao = [g for g in grupos.values() if g["missao"] is not None]
        sem_missao = [g for g in grupos.values() if g["missao"] is None]

        com_missao.sort(
            key=lambda g: max(c.criado_em for c in g["checklists"]),
            reverse=True,
        )

        ctx["grupos"] = com_missao + sem_missao
        ctx["total_geral"] = len(checklists)
        return ctx


class ChecklistDetailView(LoginRequiredMixin, DetailView):
    model = Checklist
    template_name = "checklist_detail.html"
    context_object_name = "checklist"

    def get_queryset(self):
        return (
            Checklist.objects
            .select_related("missao")
            .prefetch_related(
                Prefetch(
                    "itens",
                    queryset=ChecklistItem.objects
                    .select_related("responsavel", "concluido_por")
                    .prefetch_related("fotos__enviado_por", "eventos__autor"),
                ),
            )
        )

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["itens"] = self.object.itens.select_related(
            "responsavel", "concluido_por"
        ).prefetch_related("fotos__enviado_por", "eventos__autor")
        ctx["import_form"] = ChecklistImportForm()
        ctx["importacoes"] = self.object.importacoes.all()[:5]

        # Mensagem pronta do WhatsApp
        url_auditoria = self.request.build_absolute_uri(
            reverse("auditoria_checklist", args=[self.object.token_auditoria])
        )
        ctx["mensagem_whatsapp"] = (
            f"*Auditoria do checklist:* {self.object.titulo}\n"
            f"{self.object.feitos}/{self.object.total} itens concluídos\n\n"
            f"Abra o link abaixo e audite:\n{url_auditoria}"
        )
        return ctx


class ChecklistCreateView(TrabalhadorRequiredMixin, CreateView):
    model = Checklist
    form_class = ChecklistForm
    template_name = "checklist_form.html"

    def form_valid(self, form):
        form.instance.criado_por = self.request.user
        messages.success(
            self.request, "Checklist criado. Agora importe a planilha."
        )
        return super().form_valid(form)


class ChecklistUpdateView(TrabalhadorRequiredMixin, UpdateView):
    model = Checklist
    form_class = ChecklistForm
    template_name = "checklist_form.html"

    def form_valid(self, form):
        messages.success(self.request, "Checklist atualizado.")
        return super().form_valid(form)


class ChecklistImportView(TrabalhadorRequiredMixin, View):
    """Recebe o Excel, faz o parse e cria os itens."""

    def post(self, request, pk):
        checklist = get_object_or_404(Checklist, pk=pk)
        form = ChecklistImportForm(request.POST, request.FILES)
        if not form.is_valid():
            for erro in form.errors.values():
                messages.error(request, erro.as_text())
            return redirect(checklist)

        arquivo = form.cleaned_data["arquivo"]
        substituir = form.cleaned_data.get("substituir")
        criar_sistemas = form.cleaned_data.get("criar_sistemas", False)

        importacao = ChecklistImportacao.objects.create(
            checklist=checklist,
            arquivo=arquivo,
            enviado_por=request.user,
        )

        try:
            resultado = ler_excel(importacao.arquivo.path)
        except Exception as exc:
            messages.error(request, f"Falha ao ler a planilha: {exc}")
            return redirect(checklist)

        if not resultado.linhas:
            importacao.erros = "\n".join(resultado.erros) or "Nenhuma linha válida."
            importacao.save()
            messages.warning(request, "Nenhum item encontrado na planilha.")
            return redirect(checklist)

        if substituir:
            checklist.itens.all().delete()

        criados = aplicar_import(
            checklist, resultado,
            criar_faltantes=criar_sistemas,
      
        )

        importacao.linhas_lidas = len(resultado.linhas)
        importacao.linhas_criadas = criados
        importacao.erros = "\n".join(resultado.erros)
        importacao.save()

        messages.success(
            request,
            f"{criados} ite{'ns' if criados != 1 else 'm'} importado"
            f"{'s' if criados != 1 else ''} de {len(resultado.linhas)} linha(s).",
        )
        if resultado.erros:
            messages.warning(
                request,
                f"{len(resultado.erros)} linha(s) com aviso — veja o histórico.",
            )
        return redirect(checklist)


class ChecklistItemToggleView(TrabalhadorRequiredMixin, View):
    """Marca/desmarca um item como concluído (admin)."""

    def post(self, request, pk):
        item = get_object_or_404(ChecklistItem, pk=pk)
        anterior = item.concluido
        item.concluido = not item.concluido
        if item.concluido:
            item.concluido_em = timezone.now()
            item.concluido_por = request.user
        else:
            item.concluido_em = None
            item.concluido_por = None
        item.save(update_fields=[
            "concluido", "concluido_em", "concluido_por",
        ])

        if item.concluido and not anterior:
            _evento(item, ChecklistItemEvento.Tipo.CONCLUIDO, autor=request.user)
        elif not item.concluido and anterior:
            _evento(item, ChecklistItemEvento.Tipo.REABERTO, autor=request.user)

        messages.success(
            request,
            f"{'Concluído' if item.concluido else 'Reaberto'}: {item.titulo}",
        )
        return redirect(item.checklist)


# =========================================================
# MODELO DE EXCEL
# =========================================================
try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    OPENPYXL_OK = True
except ImportError:
    OPENPYXL_OK = False


class ChecklistModeloView(LoginRequiredMixin, View):
    """Gera e devolve um .xlsx de exemplo."""

    def get(self, request):
        if not OPENPYXL_OK:
            messages.error(
                request,
                "Instale o pacote 'openpyxl' para baixar o modelo "
                "(pip install openpyxl).",
            )
            return redirect("checklist_lista")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Checklist"

        cabecalhos = [
            "Ordem", "Título", "Descrição",
            "Sistema", "Prioridade", "Responsável",
        ]
        ws.append(cabecalhos)

        exemplos = [
            [1, "Checar casco", "Inspeção visual externa", "Casco", "Alta", "joao"],
            [2, "Testar motor", "Partida a frio + rpm", "Motor", "Crítica", "maria"],
            [3, "Conferir bateria", "Nível mínimo de 80%", "Bateria", "Média", ""],
            [4, "Verificar Starlink", "Sinal, ping e estabilidade", "Starlink", "Alta", "joao"],
            [5, "Checar leme", "Curso completo do leme", "Leme", "Média", ""],
        ]
        for linha in exemplos:
            ws.append(linha)

        header_fill = PatternFill("solid", fgColor="1F3A5F")
        header_font = Font(bold=True, color="FFFFFF")
        for col_idx, _ in enumerate(cabecalhos, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for i, largura in enumerate([8, 28, 36, 14, 12, 16], start=1):
            ws.column_dimensions[get_column_letter(i)].width = largura

        ws.freeze_panes = "A2"

        guia = wb.create_sheet("Instruções")
        instrucoes = [
            ["COMO PREENCHER ESTA PLANILHA"],
            [""],
            ["1. Não mude os nomes das colunas da aba 'Checklist'."],
            ["2. Só o campo 'Título' é obrigatório. Linhas sem título são ignoradas."],
            ["3. 'Ordem' é opcional — se vazio, o sistema numera automaticamente."],
            ["4. 'Prioridade' aceita: Baixa, Média, Alta, Crítica (ou Low/Medium/High/Critical)."],
            ["5. 'Sistema' é texto livre (ex.: Casco, Motor, Bateria)."],
            ["6. 'Responsável' deve ser o USERNAME do usuário (não o nome completo)."],
            ["7. Salve como .xlsx (Excel 2007+) antes de subir."],
            ["8. Só a primeira aba é lida."],
            [""],
            ["Apague as linhas de exemplo e coloque as suas."],
        ]
        for linha in instrucoes:
            guia.append(linha)
        guia.column_dimensions["A"].width = 80
        guia["A1"].font = Font(bold=True, size=13)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        resposta = HttpResponse(
            buffer.read(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )
        resposta["Content-Disposition"] = (
            'attachment; filename="modelo_checklist_l42.xlsx"'
        )
        return resposta


# =========================================================
# AUDITORIA (link público / com login opcional)
# =========================================================
from django.contrib.auth.mixins import AccessMixin


class AuditoriaAccessMixin(AccessMixin):
    """Exige login só se o checklist tiver exige_login_auditoria=True."""

    def dispatch(self, request, *args, **kwargs):
        token = kwargs.get("token")
        checklist = get_object_or_404(Checklist, token_auditoria=token)

        if checklist.exige_login_auditoria and not request.user.is_authenticated:
            return self.handle_no_permission()
        return super().dispatch(request, *args, **kwargs)


class AuditoriaChecklistView(AuditoriaAccessMixin, DetailView):
    model = Checklist
    template_name = "auditoria/checklist.html"
    context_object_name = "checklist"
    slug_field = "token_auditoria"
    slug_url_kwarg = "token"

    def get_object(self, queryset=None):
        token = self.kwargs["token"]
        return get_object_or_404(
            Checklist.objects.prefetch_related(
                "itens__fotos__enviado_por",
                "itens__concluido_por",
                "itens__eventos__autor",
            ),
            token_auditoria=token,
        )

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        ctx["itens"] = (
            self.object.itens
            .select_related("responsavel", "concluido_por")
            .prefetch_related("fotos__enviado_por", "eventos__autor")
        )
        return ctx


class AuditoriaToggleItemView(AuditoriaAccessMixin, View):
    def post(self, request, token, item_id):
        checklist = get_object_or_404(Checklist, token_auditoria=token)
        item = get_object_or_404(ChecklistItem, pk=item_id, checklist=checklist)

        anterior = item.concluido
        item.concluido = not item.concluido

        if item.concluido:
            item.concluido_em = timezone.now()
            item.concluido_por = (
                request.user if request.user.is_authenticated else None
            )
        else:
            item.concluido_em = None
            item.concluido_por = None

        obs = request.POST.get("observacao", "").strip()
        if obs and obs != item.observacao:
            item.observacao = obs
            _evento(
                item, ChecklistItemEvento.Tipo.OBSERVACAO,
                autor=request.user if request.user.is_authenticated else None,
                texto=obs,
            )

        item.save(update_fields=[
            "concluido", "concluido_em", "concluido_por", "observacao",
        ])

        autor = request.user if request.user.is_authenticated else None
        if item.concluido and not anterior:
            _evento(item, ChecklistItemEvento.Tipo.CONCLUIDO, autor=autor)
        elif not item.concluido and anterior:
            _evento(item, ChecklistItemEvento.Tipo.REABERTO, autor=autor)

        nome = (
            request.user.get_full_name() or request.user.username
            if request.user.is_authenticated else "anônimo"
        )
        messages.success(
            request,
            f"{'Concluído' if item.concluido else 'Reaberto'} por {nome}.",
        )
        return redirect("auditoria_checklist", token=token)


class AuditoriaUploadFotoView(AuditoriaAccessMixin, View):
    def post(self, request, token, item_id):
        checklist = get_object_or_404(Checklist, token_auditoria=token)
        item = get_object_or_404(ChecklistItem, pk=item_id, checklist=checklist)

        imagem = request.FILES.get("imagem")
        if not imagem:
            messages.error(request, "Selecione uma imagem.")
            return redirect("auditoria_checklist", token=token)

        if not imagem.content_type.startswith("image/"):
            messages.error(request, "O arquivo precisa ser uma imagem.")
            return redirect("auditoria_checklist", token=token)

        legenda = request.POST.get("legenda", "").strip()
        autor = request.user if request.user.is_authenticated else None

        foto = ChecklistItemFoto.objects.create(
            item=item, imagem=imagem, legenda=legenda,
            enviado_por=autor,
        )

        _evento(
            item, ChecklistItemEvento.Tipo.FOTO, autor=autor,
            texto=legenda or f"Foto enviada ({foto.imagem.name.split('/')[-1]})",
        )

        if request.POST.get("concluir") == "on" and not item.concluido:
            item.concluido = True
            item.concluido_em = timezone.now()
            item.concluido_por = autor
            item.save(update_fields=["concluido", "concluido_em", "concluido_por"])
            _evento(item, ChecklistItemEvento.Tipo.CONCLUIDO, autor=autor)

        nome = (
            request.user.get_full_name() or request.user.username
            if request.user.is_authenticated else "anônimo"
        )
        messages.success(request, f"Foto anexada por {nome}.")
        return redirect("auditoria_checklist", token=token)


class AuditoriaRemoverFotoView(AuditoriaAccessMixin, View):
    def post(self, request, token, foto_id):
        checklist = get_object_or_404(Checklist, token_auditoria=token)
        foto = get_object_or_404(
            ChecklistItemFoto, pk=foto_id, item__checklist=checklist
        )

        autor = request.user if request.user.is_authenticated else None

        if (
            foto.enviado_por
            and foto.enviado_por != request.user
            and not request.user.is_superuser
        ):
            messages.error(request, "Você só pode remover fotos que você enviou.")
            return redirect("auditoria_checklist", token=token)

        item = foto.item
        _evento(
            item, ChecklistItemEvento.Tipo.FOTO_REMOVIDA,
            autor=autor, texto=foto.legenda or "",
        )

        foto.imagem.delete(save=False)
        foto.delete()
        messages.success(request, "Foto removida.")
        return redirect("auditoria_checklist", token=token)

class MissaoDetailView(LoginRequiredMixin, DetailView):
    model = Missao
    template_name = "missao_detail.html"
    context_object_name = "missao"

    def get_queryset(self):
        return (
            Missao.objects
            .select_related("responsavel", "criado_por")
            .prefetch_related(
                "checklists__itens__responsavel",
                "checklists__itens__fotos",
                "checklists__itens__eventos",
            )
        )

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        m = self.object

        # Progresso agregado dos checklists
        checklists = list(m.checklists.all())
        if checklists:
            m.total_itens = sum(c.itens.count() for c in checklists)
            m.itens_feitos = sum(
                c.itens.filter(concluido=True).count() for c in checklists
            )
            m.progresso_geral = (
                m.itens_feitos * 100 // m.total_itens if m.total_itens else 0
            )
        else:
            m.total_itens = 0
            m.itens_feitos = 0
            m.progresso_geral = 0

        # Demandas recentes ligadas à mesma embarcação/sistema
        # (aqui não há FK direta de Demanda -> Missao, então mostramos
        #  as mais recentes em aberto para dar contexto operacional)
        ctx["demandas_recentes"] = (
            Demanda.objects
            .filter(status__in=STATUS_ABERTOS)
            .select_related("sistema", "responsavel")
            .ordenadas()[:6]
        )

        ctx["checklists"] = checklists
        return ctx